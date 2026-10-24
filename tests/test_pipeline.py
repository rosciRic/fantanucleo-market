import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import openpyxl
import organizza_fanta as importer
from export_json import generate_giornata_json, build_day, read_tables, tables_from_data
from site_data import ROOT, check, validate, manifest


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.data = json.loads((ROOT / 'giornata_7/data_7.json').read_text())

    def test_all_committed_days_are_consistent(self):
        self.assertTrue(check()['giornate'])

    def test_missing_roster_member_is_rejected(self):
        next(iter(self.data['rose'].values())).pop()
        with self.assertRaises(ValueError):
            validate(self.data, 7)

    def test_wrong_owner_is_rejected(self):
        self.data['giocatori'][0]['owners'] = ['Squadra inesistente']
        with self.assertRaises(ValueError):
            validate(self.data, 7)

    def test_invalid_market_balances_are_rejected(self):
        for field in ('cambi', 'crediti'):
            for value in ('—', None, float('nan'), -1, True):
                with self.subTest(field=field, value=value):
                    data = json.loads((ROOT / 'giornata_7/data_7.json').read_text())
                    data['cambi'][0][field] = value
                    with self.assertRaises(ValueError):
                        validate(data, 7)

    def test_manifest_is_deterministic(self):
        self.assertEqual(manifest(), manifest())

    def test_json_results_roundtrip_preserves_legacy_archive(self):
        for source in ROOT.glob('giornata_*/quotazioni_*.csv'):
            day = int(source.parent.name.split('_')[1])
            with self.subTest(day=day):
                folder = source.parent
                expected = generate_giornata_json(folder, day)
                tables = read_tables(folder, day)
                recovered = tables_from_data(expected)
                recovered['quotazioni'] = tables['quotazioni']
                self.assertEqual(build_day(recovered, day), expected)

    def test_changing_one_day_keeps_other_cache_versions(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for day in (1, 2):
                folder = root / f'giornata_{day}'
                folder.mkdir()
                self.data['giornata'] = day
                (folder / f'data_{day}.json').write_text(json.dumps(self.data))
            before = manifest(root)
            self.data['standings'][0]['pts'] += 1
            (root / 'giornata_2/data_2.json').write_text(json.dumps(self.data))
            after = manifest(root)
            self.assertEqual(before['versions']['1'], after['versions']['1'])
            self.assertNotEqual(before['versions']['2'], after['versions']['2'])
            self.assertNotEqual(before['version'], after['version'])

    def test_quotes_update_works_without_csv_and_exports_only_on_request(self):
        self.data = generate_giornata_json(ROOT / 'giornata_7', 7)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            folder = root / 'giornata_7'
            folder.mkdir()
            (folder / 'data_7.json').write_text(json.dumps(self.data))
            quotes = root / 'Quotazioni.xlsx'
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(['header'])
            ws.append(['header'])
            for index, row in enumerate(read_tables(ROOT / 'giornata_7', 7)['quotazioni'], 1):
                ws.append([index, row['Ruolo'], '', row['Nome'], row['Squadra'], int(row['QtA']), int(row['QtI']), int(row['Diff'])])
            wb.save(quotes)
            wb.close()
            with patch.object(importer, 'FANTA_DIR', root):
                importer.import_day(None, quotes, day=7)
                self.assertFalse(list(folder.glob('*.csv')))
                self.assertEqual(json.loads((folder / 'data_7.json').read_text()), self.data)
                importer.import_day(None, quotes, day=7, export_csv=True)
                self.assertEqual(len(list(folder.glob('*.csv'))), 5)
                self.assertEqual(generate_giornata_json(folder, 7), self.data)

    def test_updated_listone_wins_over_old_free_agent_price(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / 'quotazioni_1.csv').write_text('Nome,Squadra,Ruolo,QtA,QtI,Diff\nMario,Roma,C,0,10,-10\n')
            (root / 'svincolati_1.csv').write_text('Nome,Squadra,Ruolo,Quotazione\nMario,Roma,C,15\n')
            self.assertEqual(generate_giornata_json(root, 1)['giocatori'][0]['quot'], 0)

    def test_day_does_not_use_season_as_matchday(self):
        wb = openpyxl.Workbook()
        self.assertEqual(importer.extract_giornata_number(Path('Risultati_7ª_2026_2027.xlsx'), wb), 7)
        with self.assertRaises(ValueError):
            importer.extract_giornata_number(Path('Risultati_2026_2027.xlsx'), wb)
        wb.close()

    def test_invalid_workbook_preserves_existing_data_and_sources(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            results, quotes = root / 'Risultati_8.xlsx', root / 'Quotazioni.xlsx'
            wb = openpyxl.Workbook()
            wb.save(results)
            ws = wb.active
            ws.append(['header'])
            ws.append(['header'])
            ws.append([1, 'P', '', 'Mario', 'Roma', 10, 10, 0])
            wb.save(quotes)
            wb.close()
            (root / 'manifest.json').write_text('{"giornate":[]}')
            before = {p.name: p.read_bytes() for p in root.iterdir()}
            with patch.object(importer, 'FANTA_DIR', root):
                with self.assertRaises(ValueError):
                    importer.import_day(results, quotes, day=8)
            self.assertEqual(before, {p.name: p.read_bytes() for p in root.iterdir()})


if __name__ == '__main__':
    unittest.main()
