#!/usr/bin/env python3
"""
organizza_fanta.py — Converte i file Risultati Nª 2026_2027.xlsx e Quotazioni nei CSV per la dashboard.

Genera i CSV essenziali per ogni giornata:
  1. classifiche_campionato_N.csv  → Classifica Generale + Punti
  2. rose_N.csv                    → Rose di tutte le 18 squadre
  3. svincolati_N.csv              → Calciatori liberi sul mercato
  4. cambi_N.csv                   → Cambi e crediti residui per ogni squadra
  5. quotazioni_N.csv              → Quotazioni attuali e variazioni dal listone ufficiale

Uso:
  python3 organizza_fanta.py                  # Importa una giornata dai Download, senza pubblicare
  python3 organizza_fanta.py percorso.xlsx    # Elabora un file specifico
"""
import sys
import re
import csv
import shutil
import tempfile
import subprocess
from zipfile import BadZipFile
from pathlib import Path
import openpyxl

DOWNLOADS_DIR = Path.home() / "Downloads"
FANTA_DIR = Path(__file__).resolve().parent


def extract_giornata_number(filepath: Path, wb: openpyxl.Workbook) -> int:
    """Ricava il numero della giornata dal nome del file o dal foglio."""
    match = re.search(r'risultati[ _-]*(\d{1,2})(?:[ªa _.-]|$)', filepath.name, re.I)
    if not match:
        match = re.search(r'(\d{1,2})\s*[ªa]\s*(?:giornata)?', str(wb.worksheets[0].cell(1, 1).value or ''), re.I)
    if not match or not 1 <= int(match.group(1)) <= 99:
        raise ValueError('Giornata non riconosciuta. Usa --day N o Risultati_Nª_2026_2027.xlsx')
    return int(match.group(1))


def cell_str(cell):
    return str(cell).strip() if cell is not None else ''


def parse_classifica(sheet):
    """Estrae la Classifica Generale dal foglio risultati."""
    rows = list(sheet.iter_rows(values_only=True))
    campionato = []
    mode = None

    for r in rows:
        line = [cell_str(c) for c in r]
        if not any(line):
            continue
        joined = ' '.join(line)

        if 'CLASSIFICA GENERALE' in joined:
            mode = 'gen'
            continue
        elif any(kw in joined for kw in ['CLASSIFICA DI GIORNATA', 'CLASSIFICA COPPA', 'CAMBI DA FARE']):
            mode = None
            continue

        if mode == 'gen' and len(line) >= 4 and line[1].isdigit():
            campionato.append([
                'Generale',
                int(line[1]),
                line[2],
                float(line[3])
            ])

    return campionato


def parse_rose(sheet):
    """Estrae la composizione delle 18 rose (23 calciatori per squadra)."""
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    RUOLI = [(0, 3, 'P'), (3, 10, 'D'), (10, 18, 'C'), (18, 23, 'A')]
    parsed = []
    i = 0

    while i < len(rows):
        r = rows[i]
        teams = []
        for c in range(0, len(r), 3):
            name = cell_str(r[c])
            if name and name.lower() not in ('costo', 'totale', ''):
                teams.append((name, c))

        if teams:
            start = i + 1
            for p_idx in range(start, min(start + 25, len(rows))):
                p_row = rows[p_idx]
                if any(cell_str(c).lower() == 'totale' for c in p_row if c is not None):
                    break
                offset = p_idx - start
                ruolo = next((r for lo, hi, r in RUOLI if lo <= offset < hi), 'C')

                for t_name, c_col in teams:
                    if c_col < len(p_row) and p_row[c_col] is not None and cell_str(p_row[c_col]):
                        nome = cell_str(p_row[c_col])
                        costo = cell_str(p_row[c_col + 1]) if c_col + 1 < len(p_row) else '0'
                        parsed.append([t_name, ruolo, nome, int(costo) if costo.isdigit() else 0])

            i = start + 23
        else:
            i += 1

    return parsed


def parse_svincolati(sheet, quot_map=None):
    """Estrae l'elenco dei calciatori svincolati liberi. I calciatori fuori lista o andati all'estero sono contrassegnati con * e quotazione 0."""
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    header = [cell_str(c).upper() for c in rows[0]]

    def find_col(patterns, default_idx):
        for p in patterns:
            for idx, h in enumerate(header):
                if p in h:
                    return idx
        return default_idx

    idx_num = find_col(['#'], 0)
    idx_nome = find_col(['NOME'], 1)
    idx_fuori = find_col(['FUORI LISTA', 'FUORI'], 2)
    idx_sq = find_col(['SQ.'], 3)
    idx_ruolo = find_col(['R.'], 5)
    idx_pg = find_col(['PGV', 'PG'], 6)
    idx_mv = find_col(['MV'], 7)
    idx_fm = find_col(['FM'], 8)
    idx_fvm = find_col(['FVM'], 9)
    idx_quot = find_col(['QUOT.'], len(header) - 1)

    parsed = []

    for r in rows[1:]:
        if not r or not any(c is not None for c in r):
            continue
        line = [cell_str(c) for c in r]

        if idx_num < len(line) and not line[idx_num].isdigit():
            continue

        nome = line[idx_nome] if idx_nome < len(line) else ''
        if not nome:
            continue

        fuori = line[idx_fuori] if idx_fuori < len(line) else ''

        # Se 'fuori' o 'nome' contiene '*', è fuori lista / andato all'estero
        is_fuori_lista = ('*' in fuori) or ('*' in nome)
        clean_nome = nome.replace('*', '').strip()

        # Quotazione dal foglio o dal quot_map (listone)
        raw_quot = line[idx_quot] if idx_quot < len(line) else '0'
        if not raw_quot or not raw_quot.isdigit():
            raw_quot = '0'

        # Preferisci la quotazione del listone se disponibile
        if quot_map and clean_nome.lower() in quot_map:
            raw_quot = str(quot_map[clean_nome.lower()])

        # Se è fuori lista o la quotazione è 0 (estero/fuori rosa)
        if is_fuori_lista or raw_quot == '0':
            display_nome = f"{clean_nome}*"
            final_quot = '0'
        else:
            display_nome = clean_nome
            final_quot = raw_quot

        sq = line[idx_sq] if idx_sq < len(line) else ''
        ruolo = line[idx_ruolo] if idx_ruolo < len(line) else ''
        pg = line[idx_pg] if idx_pg < len(line) else '0'
        mv = line[idx_mv] if idx_mv < len(line) else '0'
        fm = line[idx_fm] if idx_fm < len(line) else '0'
        fvm = line[idx_fvm] if idx_fvm < len(line) else '0'

        parsed.append([
            display_nome,
            sq,
            ruolo,
            pg,
            mv,
            fm,
            fvm,
            final_quot,
        ])

    return parsed


def parse_cambi(sheet):
    """Estrae la tabella cambi e crediti dal foglio risultati."""
    rows_data = list(sheet.iter_rows(values_only=True))
    start = None
    for i, row in enumerate(rows_data):
        if any('CAMBI DA FARE' in str(c) for c in row if c):
            start = i + 1
            break
    if start is None:
        return []

    result = []
    for row in rows_data[start:]:
        vals = [c for c in row if c is not None]
        if not vals:
            break
        non_none = [(i, row[i]) for i in range(len(row)) if row[i] is not None]
        if len(non_none) < 3:
            break
        nome = str(non_none[0][1]).strip()
        cambi = non_none[1][1] if len(non_none) > 1 else ''
        crediti = non_none[2][1] if len(non_none) > 2 else ''
        extra = str(non_none[3][1]).strip() if len(non_none) > 3 else ''

        if not nome or nome.upper() in ('FANTASQUADRA', 'SQUADRA'):
            continue
        result.append([nome, cambi, crediti, extra])

    return result


def parse_listone_file(source: Path, num: int):
    """Legge soltanto il listone selezionato per questa importazione."""
    file_path = source
    wb = None
    print(f"  → Lettura listone per G{num} da: {file_path.name}")
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb['Tutti'] if 'Tutti' in wb.sheetnames else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        parsed = []
        for r in rows[2:]:
            if len(r) >= 8 and r[3] is not None:
                nome = str(r[3]).strip()
                sq = str(r[4]).strip() if len(r) > 4 else ''
                ruolo = str(r[1]).strip() if len(r) > 1 else ''
                qti = str(r[6]).strip() if len(r) > 6 else '0'
                diff = str(r[7]).strip() if len(r) > 7 else '0'
                
                # La prima giornata usa Qt.I; zero e valori mancanti sono distinti.
                selected = r[6] if num == 1 else r[5]
                fallback = r[5] if num == 1 else r[6]
                valore = str(selected if selected is not None else fallback).strip()
                int(valore)  # Rifiuta quotazioni mancanti o non numeriche.
                parsed.append([nome, sq, ruolo, valore, qti, diff])
        return parsed
    except Exception as e:
        print(f"Errore lettura listone ({file_path.name}): {e}")
        return []
    finally:
        if wb:
            wb.close()


def write_csv(path, header, rows):
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    return len(rows)


# Import, preview and publication are separate operations.
def publish(day):
    folder = f'giornata_{day}'
    paths = ['manifest.json'] + [str(p.relative_to(FANTA_DIR)) for p in sorted((FANTA_DIR / folder).iterdir()) if p.suffix in ('.json', '.csv')]
    def git(*args, capture=False):
        return subprocess.run(['git', *args], cwd=FANTA_DIR, check=True, text=True, capture_output=capture)
    branch = git('branch', '--show-current', capture=True).stdout.strip()
    if branch != 'main':
        raise ValueError('Pubblicazione consentita da main. I dati locali sono pronti.')
    staged = git('diff', '--cached', '--name-only', capture=True).stdout.splitlines()
    if staged:
        raise ValueError('Ci sono file già in stage: completa quel commit prima di pubblicare i dati.')
    git('add', '--', *paths)
    changed = git('diff', '--cached', '--name-only', capture=True).stdout.strip()
    if changed:
        git('commit', '-m', f'Aggiornamento dati Giornata {day}', '--only', '--', *paths)
    git('push', 'origin', 'main')
    print('Push completato. Controlla il deploy nella scheda Actions di GitHub.')


def import_day(results, quotes, day=None, dry_run=False):
    from export_json import generate_giornata_json
    from site_data import validate, atomic_json, manifest, days
    wb = None
    try:
        if results:
            wb = openpyxl.load_workbook(results, data_only=True)
            day = day or extract_giornata_number(results, wb)
        else:
            available = days(FANTA_DIR)
            if not available:
                raise ValueError('Importa prima un file risultati.')
            day = day or available[-1]
        output = FANTA_DIR / f'giornata_{day}'
        if not results and not output.is_dir():
            raise ValueError('La giornata da aggiornare non esiste.')
        if quotes is None:
            archived = list(output.glob('Quotazioni*.xlsx'))
            if len(archived) == 1:
                quotes = archived[0]
        if quotes is None:
            raise ValueError('Serve un listone quotazioni esplicito o già archiviato per questa giornata.')
        with tempfile.TemporaryDirectory(prefix='.import_', dir=FANTA_DIR) as temp:
            staging = Path(temp)
            if output.exists():
                for source in output.glob('*.csv'):
                    shutil.copy2(source, staging / source.name)
            values = parse_listone_file(quotes, day)
            if not values:
                raise ValueError('Listone vuoto o non riconosciuto.')
            write_csv(staging / f'quotazioni_{day}.csv', ['Nome','Squadra','Ruolo','QtA','QtI','Diff'], values)
            if wb:
                def sheet(part):
                    names = [s for s in wb.sheetnames if part in s.lower()]
                    if len(names) != 1:
                        raise ValueError(f'Foglio {part} mancante o ambiguo.')
                    return wb[names[0]]
                results_sheet = sheet('risultat')
                tables = [
                    ('classifiche_campionato', ['Tipo','Posizione','Fantasquadra','Punti'], parse_classifica(results_sheet)),
                    ('rose', ['Fantasquadra','Ruolo','Calciatore','Costo'], parse_rose(sheet('rose'))),
                    ('cambi', ['Fantasquadra','CambiRimasti','CreditiRimasti','CambiExtra'], parse_cambi(results_sheet)),
                    ('svincolati', ['Nome','Squadra','Ruolo','PG','MV','FM','FVM','Quotazione'], parse_svincolati(sheet('svincolat'), {q[0].replace('*','').strip().lower(): q[3] for q in values})),
                ]
                for name, header, rows in tables:
                    write_csv(staging / f'{name}_{day}.csv', header, rows)
            data = generate_giornata_json(staging, day)
            validate(data, day)
            atomic_json(staging / f'data_{day}.json', data)
            print(f'G{day}: {len(data["giocatori"])} giocatori, {len(data["rose"])} rose. Dati validi.')
            if dry_run:
                print('Simulazione completata: nessun dato o file sorgente modificato.')
                return day
            # Keep the original downloads. Copy only the exact files used in this import.
            shutil.copy2(quotes, staging / f'Quotazioni_{day}ª_2026_2027.xlsx')
            if results:
                shutil.copy2(results, staging / f'Risultati_{day}ª_2026_2027.xlsx')
            backup = Path(temp) / 'previous'
            if output.exists():
                shutil.copytree(output, backup)
            output.mkdir(exist_ok=True)
            try:
                for source in staging.iterdir():
                    if source.is_file():
                        source.replace(output / source.name)
                atomic_json(FANTA_DIR / 'manifest.json', manifest(FANTA_DIR))
            except Exception:
                shutil.rmtree(output)
                if backup.exists():
                    shutil.copytree(backup, output)
                raise
            print(f'Dati salvati in {output}.')
            return day
    finally:
        if wb:
            wb.close()


def main():
    import argparse
    parser = argparse.ArgumentParser(description='Importa e valida una giornata Fantanucleo. Pubblicazione solo con --publish.')
    parser.add_argument('results', nargs='?', type=Path, help='Excel risultati (facoltativo per aggiornare solo le quotazioni)')
    parser.add_argument('--quotes', type=Path, help='Excel quotazioni da associare')
    parser.add_argument('--day', type=int, choices=range(1, 100), metavar='N')
    parser.add_argument('--downloads', type=Path, default=DOWNLOADS_DIR)
    action = parser.add_mutually_exclusive_group()
    action.add_argument('--dry-run', action='store_true', help='Valida senza modificare i dati')
    action.add_argument('--publish', action='store_true', help='Importa, committa i dati e fa push su main')
    args = parser.parse_args()
    try:
        results, quotes = args.results, args.quotes
        # Explicit inputs bypass discovery of unrelated results in Downloads.
        if results is None and quotes is None:
            candidates = [p for p in args.downloads.glob('*.xlsx') if 'risultati' in p.name.lower()]
            if len(candidates) > 1:
                raise ValueError('Più file risultati trovati: indica esplicitamente quello da importare.')
            results = candidates[0] if candidates else None
        if quotes is None:
            candidates = list(args.downloads.glob('Quotazioni_Fantacalcio*.xlsx'))
            if len(candidates) > 1:
                raise ValueError('Più listoni trovati: selezionane uno con --quotes.')
            quotes = candidates[0] if candidates else None
        if results is None and quotes is None:
            raise ValueError('Nessun Excel trovato. Indica i percorsi dei file o usa --downloads.')
        for path in (results, quotes):
            if path and not path.is_file():
                raise ValueError(f'File non trovato: {path}')
        day = import_day(results, quotes, args.day, args.dry_run)
        if args.publish:
            publish(day)
        elif not args.dry_run:
            print('Anteprima: python3 -m http.server 8000 --bind 127.0.0.1. Aggiungi --publish per pubblicare.')
        return 0
    except (ValueError, OSError, KeyError, BadZipFile, openpyxl.utils.exceptions.InvalidFileException, subprocess.CalledProcessError) as error:
        print(f'Errore: {error}', file=sys.stderr)
        return 1


if __name__ == '__main__':
    sys.exit(main())
